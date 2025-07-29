from django.shortcuts import render
from datetime import timedelta, date, time
from .models import *
from django.contrib.auth.hashers import check_password
from django.db import transaction
from django.utils import timezone
from decimal import Decimal, ROUND_HALF_UP
from rest_framework.generics import ListCreateAPIView, RetrieveUpdateDestroyAPIView,ListAPIView
import csv
import os
from rest_framework.response import Response
import re
import base64
from datetime import datetime
from django.http import Http404
import uuid
from weasyprint import HTML, CSS
from django.template.loader import render_to_string
from rest_framework.decorators import api_view
#from django.views.decorators.csrf import csrf_exempt
#from django.utils.decorators import method_decorator
from django.http.response import JsonResponse, HttpResponse
from django.db.models import Sum, Count, F, DecimalField, Q
from django.db.models.functions import Coalesce
from rest_framework.permissions import AllowAny, IsAuthenticated
from http import HTTPStatus
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.views import APIView
from rest_framework import status
from rest_framework_simplejwt.views import TokenObtainPairView
from rest_framework_simplejwt.serializers import TokenObtainPairSerializer
# --- LIBRERÍAS PARA PDF Y EXCEL ---
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, landscape
from openpyxl.drawing.image import Image
from reportlab.lib.units import inch
from reportlab.lib import colors
from openpyxl import Workbook
from .woocommerce_sync import actualizar_stock_woocommerce
from openpyxl.styles import Font, Alignment
from rest_framework import generics, permissions
from .serializers import (
    ProductoSerializer,
    ClienteSerializer,
    CategoriaSerializer,
    AlmacenSerializer,
    InventarioSerializer,
    UserMetadataSerializer,
    ConfiguracionivaSerializer,
    CupondescuentoSerializer,
    DetallefacturaSerializer,
    DevolucionSerializer,
    FacturaSerializer,
    FacturaelectronicaSerializer,
    MovimientoInventarioSerializer,
    LecturacodigobarrasSerializer,
    LogactividadSerializer,
    OrdenSerializer,
    ProductocategoriaSerializer,
    ProveedorSerializer,
    ReporteclienteSerializer,
    ReporteinventarioSerializer,
    ReservastockSerializer,
    RolSerializer,
    SesionusuarioSerializer,
    TipodocumentofiscalSerializer,
    TransaccionpagoSerializer,
    VariacionproductoSerializer,
    ConfiguracionCorrelativoReadSerializer,
    ConfiguracionCorrelativoWriteSerializer,
    MetodoPagoSerializer,AtributoSerializer,
    ValorAtributoSerializer, AtributoCrearConValoresSerializer, 
    UserMetadataCreateSerializer, MyTokenObtainPairSerializer, 
    FacturaReporteSerializer, UserDetailSerializer, AsistenciaSerializer,
    HorarioSerializer, DiaFestivoSerializer,
    FacturaReportSerializer, VentaReporteSerializer,
    InventarioItemSerializer
)
from django.contrib.auth import authenticate, get_user_model
from rest_framework.permissions import BasePermission

User = get_user_model()  # Esto utiliza el modelo de usuario de Django.


# ------------------------------------------------------------------------------
# Registro de usuario minimal (solo correo, password y rol)
# ------------------------------------------------------------------------------

class IsAdmin(BasePermission):
    def has_permission(self, request, view):
        if not request.user or not request.user.is_authenticated:
            return False
        metadata = getattr(request.user, 'metadata', None)
        return metadata and metadata.rol and metadata.rol.nombre == 'Administrador'

class MyTokenObtainPairView(TokenObtainPairView):
     permission_classes = [AllowAny]
     serializer_class = MyTokenObtainPairSerializer


class DashboardDataView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):
        # --- 1. FILTRO DE FECHAS ---
        end_date = timezone.now().date()
        start_date_str = request.query_params.get('start_date', (end_date - timedelta(days=29)).strftime('%Y-%m-%d'))
        end_date_str = request.query_params.get('end_date', end_date.strftime('%Y-%m-%d'))

        base_queryset = Factura.objects.filter(fecha_operacion__date__range=[start_date_str, end_date_str])
        ventas_queryset = base_queryset.exclude(estado__iexact='cancelada')

        # --- 2. CÁLCULO PARA LAS TARJETAS DE RESUMEN (Tu código ya era correcto aquí) ---
        resumen_ventas = ventas_queryset.aggregate(
            total_vendido=Coalesce(Sum('total'), Decimal('0.0')),
            total_pagado=Coalesce(Sum('total', filter=Q(estado__iexact='pagado')), Decimal('0.0')),
            total_pendiente=Coalesce(Sum('total', filter=Q(estado__iexact='pendiente')), Decimal('0.0')),
            num_transacciones=Count('id')
        )

        # --- 3. CÁLCULO PARA EL GRÁFICO DE VENTAS DIARIAS (CORREGIDO Y MEJORADO) ---
        ventas_por_dia_qs = ventas_queryset.values('fecha_operacion__date').annotate(
            total=Coalesce(Sum('total'), Decimal('0.0'))
        ).order_by('fecha_operacion__date')

        # Creamos un diccionario para mapear las ventas por fecha
        sales_map = {item['fecha_operacion__date']: item['total'] for item in ventas_por_dia_qs}

        # Generamos un rango completo de fechas para el gráfico
        start_date_obj = date.fromisoformat(start_date_str)
        end_date_obj = date.fromisoformat(end_date_str)
        all_dates = [start_date_obj + timedelta(days=i) for i in range((end_date_obj - start_date_obj).days + 1)]

        grafico_ventas = {
            'labels': [d.strftime('%d/%m') for d in all_dates],
            'data': [sales_map.get(d, Decimal('0.0')) for d in all_dates] # Usamos 0 para días sin ventas
        }

        # --- 4. PRODUCTOS MÁS VENDIDOS (CORREGIDO) ---
        productos_mas_vendidos_qs = Detallefactura.objects.filter(factura__in=ventas_queryset) \
            .values('producto__nombre', 'variante__nombre', 'producto__imagen') \
            .annotate(
                cantidad_total=Coalesce(Sum('cantidad'), 0) # Usamos Coalesce para evitar nulls
            ).order_by('-cantidad_total')[:5]
        
        productos_mas_vendidos = []
        for p in productos_mas_vendidos_qs:
            imagen_url = None
            if p.get('producto__imagen'):
                imagen_url = request.build_absolute_uri(settings.MEDIA_URL + p['producto__imagen'])
            
            productos_mas_vendidos.append({
                'nombre': f"{p['producto__nombre']} ({p['variante__nombre']})" if p['variante__nombre'] else p['producto__nombre'],
                'cantidad': p['cantidad_total'],
                'imagen': imagen_url
            })

        # --- 5. PRODUCTOS CON BAJO STOCK (Tu código ya era correcto aquí) ---
        productos_bajo_stock_qs = Producto.objects.filter(cantidad__lt=10, activo=True).order_by('cantidad')[:5]
        productos_bajo_stock = []
        for p in productos_bajo_stock_qs:
            imagen_url = p.imagen.url if p.imagen else None
            productos_bajo_stock.append({'nombre': p.nombre, 'cantidad': p.cantidad or 0, 'imagen': imagen_url})
        
        # --- 6. TOP CATEGORÍAS (CORREGIDO) ---
        top_categorias_qs = Detallefactura.objects.filter(factura__in=ventas_queryset) \
            .values('producto__categoria__nombre') \
            .annotate(
                total_vendido=Coalesce(Sum('total_linea'), Decimal('0.0')) # Usamos Coalesce
            ) \
            .order_by('-total_vendido')[:3]
        
        top_categorias = {
            'labels': [c['producto__categoria__nombre'] for c in top_categorias_qs if c['producto__categoria__nombre']],
            'data': [float(c['total_vendido']) for c in top_categorias_qs if c['producto__categoria__nombre']]
        }
        
        # --- 7. INFORMACIÓN GENERAL (Tu código ya era correcto aquí) ---
        info_general = {
            'clientes': Cliente.objects.filter(activo=True).count(),
            'productos': Producto.objects.filter(activo=True).count(),
            'ordenes_periodo': ventas_queryset.count()
        }
        
        # --- 8. DATOS DEL USUARIO (Tu código ya era correcto aquí) ---
        user_nombre = "Admin"
        if request.user.is_authenticated:
            user_nombre = request.user.get_full_name() or request.user.username

        # --- 9. ENSAMBLAR RESPUESTA FINAL ---
        data = {
            'userInfo': {'nombre': user_nombre},
            'resumen': resumen_ventas,
            'graficoVentas': grafico_ventas,
            'productosMasVendidos': productos_mas_vendidos,
            'productosBajoStock': productos_bajo_stock,
            'topCategorias': top_categorias,
            'infoGeneral': info_general,
        }
        return JsonResponse(data)



class UserMetadataListCreateView(generics.ListCreateAPIView):
    permission_classes = [permissions.IsAdminUser] # Solo los admins pueden ver y crear
    
    def get_serializer_class(self):
        # Usa el serializer de creación para peticiones POST
        if self.request.method == 'POST':
            return UserMetadataCreateSerializer
        # Usa el serializer de visualización para peticiones GET
        return UserMetadataSerializer

    def get_queryset(self):
        # Filtra por el estado 'is_active' del usuario relacionado
        queryset = UserMetadata.objects.select_related('user', 'rol','almacen_asignado').all()
        estado = self.request.query_params.get('es_activo')
        if estado is not None:
            es_activo_bool = estado.lower() in ['true', '1']
            queryset = queryset.filter(user__is_active=es_activo_bool)
        return queryset

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = serializer.save()
        
        # Devuelve la respuesta usando el serializer de visualización para mostrar el objeto completo
        read_serializer = UserMetadataSerializer(instance, context={'request': request})
        return Response(read_serializer.data, status=status.HTTP_201_CREATED)

class UserMetadataDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = UserMetadata.objects.all()
    serializer_class = UserMetadataSerializer
    permission_classes = [permissions.IsAdminUser]

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        # Ahora comprueba el estado en el modelo User
        if instance.user.is_active:
            return Response(
                {"detail": "No puedes eliminar un perfil de un usuario activo. Inactívalo primero."},
                status=status.HTTP_400_BAD_REQUEST
            )
        # Si el usuario no está activo, se puede eliminar el perfil
        return super().destroy(request, *args, **kwargs)

# --- Vistas de Gestión para Administradores ---
class HorarioDetailView(APIView):
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]

    def get(self, request, pk=1):
        # Intenta obtener el horario con pk=1, si no existe, devuelve 404
        try:
            horario = Horario.objects.get(pk=pk)
            serializer = HorarioSerializer(horario)
            return Response(serializer.data)
        except Horario.DoesNotExist:
            return Response({"detail": "Not found."}, status=status.HTTP_404_NOT_FOUND)

    def put(self, request, pk=1):
        # Usa update_or_create para actualizar el horario con pk=1 o crearlo si no existe
        horario, created = Horario.objects.update_or_create(
            pk=pk,
            defaults=request.data
        )
        serializer = HorarioSerializer(horario)
        # Decide el código de estado: 201 si fue creado, 200 si fue actualizado
        status_code = status.HTTP_201_CREATED if created else status.HTTP_200_OK
        return Response(serializer.data, status=status_code)

# Vista para LISTAR y CREAR días festivos.
class DiaFestivoListView(generics.ListCreateAPIView):
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]
    queryset = DiaFestivo.objects.all().order_by('fecha') # Devuelve ordenado por fecha
    serializer_class = DiaFestivoSerializer

# Vista para ELIMINAR un día festivo específico.
class DiaFestivoDetailView(generics.DestroyAPIView):
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]
    queryset = DiaFestivo.objects.all()
    serializer_class = DiaFestivoSerializer

# --- Vistas de Asistencia (Modificadas) ---
class UserMeView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        serializer = UserDetailSerializer(request.user)
        return Response(serializer.data)

class AttendanceSummaryView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        today = timezone.localdate()
        start_of_month = today.replace(day=1)
        
        asistencia_hoy = Asistencia.objects.filter(usuario=request.user, fecha=today).first()
        current_status = asistencia_hoy.estado_actual if asistencia_hoy else 'out'

        asistencias_mes = Asistencia.objects.filter(usuario=request.user, fecha__range=[start_of_month, today])
        festivos_mes = DiaFestivo.objects.filter(fecha__range=[start_of_month, today]).count()

        summary = {
            "total_working_days": 22,
            "absent_days": asistencias_mes.filter(estado='Ausente').count(),
            "present_days": asistencias_mes.filter(estado='Presente').count(),
            "half_days": asistencias_mes.filter(estado='Medio Día').count(),
            "late_days": asistencias_mes.filter(llegada_tarde=True).count(),
            "holidays": festivos_mes,
        }

        log_serializer = AsistenciaSerializer(asistencias_mes, many=True)
        data = {"current_status": current_status, "summary": summary, "log": log_serializer.data}
        return Response(data)

class ClockInView(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        now = timezone.now()
        local_date = timezone.localdate()
        horario = Horario.objects.first()
        es_tarde = False
        if horario:
            hora_entrada_limite = datetime.combine(now.date(), horario.hora_entrada_oficial)
            hora_entrada_limite = timezone.make_aware(hora_entrada_limite)
            limite_con_margen = hora_entrada_limite + timedelta(minutes=horario.margen_tardanza_minutos)
            if now > limite_con_margen:
                es_tarde = True

        asistencia, created = Asistencia.objects.get_or_create(
            usuario=request.user,
            fecha=local_date,
            defaults={'hora_entrada': now, 'estado_actual': 'in', 'llegada_tarde': es_tarde}
        )
        if not created:
            asistencia.hora_entrada = now
            asistencia.estado_actual = 'in'
            asistencia.llegada_tarde = es_tarde
            asistencia.save()
        return Response(status=status.HTTP_200_OK)

class ClockOutView(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        Asistencia.objects.filter(usuario=request.user, fecha=timezone.localdate()).update(
            hora_salida=timezone.now(),
            estado_actual='out'
        )
        return Response(status=status.HTTP_200_OK)

class BreakView(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        start_break = request.data.get('start', True)
        asistencia = Asistencia.objects.get(usuario=request.user, fecha=timezone.localdate())
        
        if start_break:
            Descanso.objects.create(asistencia=asistencia, inicio_descanso=timezone.now())
            asistencia.estado_actual = 'break'
        else:
            ultimo_descanso = asistencia.descansos.filter(fin_descanso__isnull=True).last()
            if ultimo_descanso:
                ultimo_descanso.fin_descanso = timezone.now()
                ultimo_descanso.save()
            asistencia.estado_actual = 'in'
        asistencia.save()
        return Response(status=status.HTTP_200_OK)


class CashClosingReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        date_str = request.query_params.get('date', None)

        if date_str:
            try:
                target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
            except ValueError:
                return Response({"error": "Formato de fecha inválido. Usar YYYY-MM-DD."}, status=400)
        else:
            # Usa la fecha actual en la zona horaria local del servidor
            target_date = timezone.localdate()

        # --- LÓGICA DE FECHA CORREGIDA ---
        # Crea el inicio y el fin del día para la fecha seleccionada
        start_of_day = datetime.combine(target_date, time.min)
        end_of_day = datetime.combine(target_date, time.max)

        # Convierte estas fechas a la zona horaria activa de Django
        # (Asegúrate de tener TIME_ZONE = 'America/Caracas' en settings.py)
        start_of_day_aware = timezone.make_aware(start_of_day)
        end_of_day_aware = timezone.make_aware(end_of_day)

        # Filtra las facturas cuyo timestamp (en UTC) esté dentro de ese rango de 24h locales
        queryset = Factura.objects.filter(
            fecha_operacion__range=(start_of_day_aware, end_of_day_aware),
            estado='pagado'  # <-- Doble chequea que este sea el estado correcto
        ).order_by('fecha_operacion')

        serializer = FacturaReportSerializer(queryset, many=True)

        data = {
            "report_date": target_date,
            "transactions": serializer.data,
        }

        return Response(data)


# Productos por id
class ProductoGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = Producto.objects.filter(id=id, activo=True)
            serializer = ProductoSerializer(data, many=True)
            return Response({"data": serializer.data}, status=HTTPStatus.OK)
        except Exception as e:
            return Response(
                {"estado": "error", "mensaje": "Productos no encontrados"},
                status=HTTPStatus.NOT_FOUND,
            )


# Listar todos los productos
class ProductoList(APIView):
    #permission_classes = [IsAuthenticated]
    def get(self, request):
        productos = Producto.objects.filter(activo=True).order_by("id")
        serializer = ProductoSerializer(productos, many=True)

        return JsonResponse({"productos": serializer.data}, status=HTTPStatus.OK)


class ProductoCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = (MultiPartParser, FormParser) 

    def post(self, request):
        # La vista ya no procesa las variantes. Solo pasa los datos al serializer.
        serializer = ProductoSerializer(data=request.data, context={'request': request})
        if serializer.is_valid():
            serializer.save()
            return JsonResponse({"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED)
        return JsonResponse({"estado": "error", "mensaje": serializer.errors}, status=HTTPStatus.BAD_REQUEST)

    def put(self, request, id):
        try:
            producto = Producto.objects.get(id=id)
        except Producto.DoesNotExist:
            return JsonResponse({"estado": "error", "mensaje": "Producto no encontrado"}, status=HTTPStatus.NOT_FOUND)
        
        # La vista ya no procesa las variantes. Solo pasa los datos al serializer.
        serializer = ProductoSerializer(producto, data=request.data, partial=True, context={'request': request})
        if serializer.is_valid():
            serializer.save()
            return JsonResponse({"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK)
        return JsonResponse({"estado": "error", "mensaje": serializer.errors}, status=HTTPStatus.BAD_REQUEST)

    def delete(self, request, pk, *args, **kwargs):
        try:
            producto = Producto.objects.get(pk=pk)
            
            # En lugar de borrar, lo desactivamos
            producto.activo = False
            producto.save()
            
            return Response({"mensaje": "Producto desactivado correctamente."}, status=status.HTTP_200_OK)
        
        except Producto.DoesNotExist:
            return Response({"error": "El producto no existe."}, status=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            # Aquí capturamos el IntegrityError si ocurriera por otra razón
            return Response({"error": f"No se pudo desactivar el producto: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)


# ==============
# Correlativo
# ==============

class ConfiguracionCorrelativoManageView(APIView):
    """
    Vista segura para ver y actualizar la configuración del correlativo.
    Solo accesible por administradores.
    """
    permission_classes = [IsAuthenticated, IsAdmin] # Solo Admins

    def get(self, request):
        """Devuelve la configuración actual usando el serializer de lectura."""
        config = obtener_configuracion_correlativo()
        # CORREGIDO: Usa el serializer que NO tiene el campo de contraseña
        serializer = ConfiguracionCorrelativoReadSerializer(config)
        return JsonResponse(serializer.data)

    def put(self, request):
        """Actualiza la configuración usando el serializer de escritura."""
        with transaction.atomic():
            config = obtener_configuracion_correlativo()
            # CORREGIDO: Usa el serializer que SÍ tiene y requiere el campo de contraseña
            serializer = ConfiguracionCorrelativoWriteSerializer(config, data=request.data, partial=True)

            if not serializer.is_valid():
                return JsonResponse(serializer.errors, status=HTTPStatus.BAD_REQUEST)

            # 1. Validar contraseña del administrador
            password = serializer.validated_data.get('password')
            if not request.user.check_password(password):
                return JsonResponse(
                    {"error": "Contraseña incorrecta. No se pudo actualizar la configuración."},
                    status=HTTPStatus.FORBIDDEN
                )

            # 2. Validar que el nuevo número no sea menor al último usado
            new_number = serializer.validated_data.get('current_number')
            if new_number is not None:
                last_factura = Factura.objects.exclude(correlativo__isnull=True).order_by('-id').first()
                if last_factura and last_factura.correlativo:
                    try:
                        last_number = int(last_factura.correlativo.split('-')[-1])
                        if new_number < last_number:
                            return JsonResponse(
                                {"error": f"El número actual no puede ser menor que el de la última factura generada ({last_number})."},
                                status=HTTPStatus.BAD_REQUEST
                            )
                    except (ValueError, IndexError):
                        pass # Ignorar si el formato del correlativo no es el esperado

            # Si todas las validaciones pasan, guardar los cambios
            serializer.save()
            return JsonResponse({"mensaje": "Configuración actualizada con éxito."})




# ==============
# Logica Escanear producto y mostrarlo en el template
# ==============
def recalcular_y_guardar_factura(factura):
    """
    Función definitiva que asume que 'detalle.precio_unitario' 
    es el PRECIO FINAL (PVP, con IVA incluido).
    """
    detalles = Detallefactura.objects.filter(factura=factura).select_related(
        'producto__configuracion_iva',
        'variante__producto__configuracion_iva'
    )
    
    total_factura_subtotal = Decimal("0.00")
    total_factura_iva = Decimal("0.00")

    detalles_a_actualizar = []
    for detalle in detalles:
        # 1. El precio unitario es el precio final con IVA.
        precio_final_unitario = detalle.precio_unitario
        descuento_porcentaje = detalle.descuento or Decimal("0.00")
        
        # 2. Aplicar el descuento al precio total de la línea.
        total_linea_con_descuento = (precio_final_unitario * detalle.cantidad) * (Decimal("1") - descuento_porcentaje / Decimal("100"))
        
        # 3. Desglosar el IVA y el subtotal a partir del total ya descontado.
        tasa_iva = Decimal("0.00")
        if detalle.producto.configuracion_iva:
            tasa_iva = Decimal(detalle.producto.configuracion_iva.porcentaje_iva) / Decimal("100")
        
        if tasa_iva > 0:
            # CÁLCULO INVERSO: Obtenemos la base dividiendo por (1 + tasa).
            subtotal_linea = (total_linea_con_descuento / (Decimal("1") + tasa_iva))
        else:
            # Si no hay IVA, el subtotal es igual al total.
            subtotal_linea = total_linea_con_descuento

        # El IVA es la diferencia entre el total y la base.
        iva_linea = total_linea_con_descuento - subtotal_linea
        
        # 4. Actualizar los campos del detalle con valores redondeados.
        detalle.subtotal_linea = subtotal_linea.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        detalle.iva_linea = iva_linea.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        detalle.total_linea = total_linea_con_descuento.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        
        detalles_a_actualizar.append(detalle)
        
        total_factura_subtotal += detalle.subtotal_linea
        total_factura_iva += detalle.iva_linea
    
    # Actualización en bloque para mayor eficiencia.
    if detalles_a_actualizar:
        Detallefactura.objects.bulk_update(detalles_a_actualizar, ['subtotal_linea', 'iva_linea', 'total_linea'])

    # Actualizar los totales de la factura.
    factura.subtotal = total_factura_subtotal
    factura.iva_total = total_factura_iva
    factura.total = total_factura_subtotal + total_factura_iva
    factura.save()


@transaction.atomic
def agg_producto_a_orden(request, cliente_id, barcode, cantidad=1, descuento_linea=None, eliminar=False):
    """
    Agrega, actualiza o elimina un producto/variante de una factura activa.
    Asigna el usuario que crea la factura por primera vez.
    """
    # 1. Búsqueda del producto (esta lógica está bien y se mantiene)
    producto_base = None
    variante_encontrada = None
    precio_final = Decimal("0.00")
    stock_disponible = 0
    
    try:
        variante_encontrada = Variacionproducto.objects.select_related('producto__configuracion_iva').get(codigo_barras=barcode)
        producto_base = variante_encontrada.producto
        precio_final = variante_encontrada.precio
        stock_disponible = variante_encontrada.cantidad or 0
    except Variacionproducto.DoesNotExist:
        try:
            producto_base = Producto.objects.select_related('configuracion_iva').get(codigo_barras=barcode, activo=True)
            precio_final = producto_base.precio
            stock_disponible = producto_base.cantidad or 0
        except Producto.DoesNotExist:
            return None, "Producto no encontrado para este código."

    factura, created = Factura.objects.get_or_create(
        cliente_id=cliente_id, 
        estado="abierta", 
        fecha_operacion__date=timezone.now().date(),
        defaults={
            'usuario': request.user, # Asigna el usuario que está logueado si se crea una nueva
            'fecha_operacion': timezone.now(), 
            'metodo_pago': None, 
            'subtotal': Decimal("0.00"), 
            'descuento_global': Decimal("0.00"), 
            'iva_total': Decimal("0.00"), 
            'total': Decimal("0.00"), 
            'activo': True
        }
    )
    
    filtro_detalle = {'factura': factura, 'producto': producto_base}
    if variante_encontrada:
        filtro_detalle['variante'] = variante_encontrada
    detalle_existente = Detallefactura.objects.filter(**filtro_detalle).first()
    
    if eliminar and detalle_existente:
        detalle_existente.delete()
    elif not eliminar:
        cantidad_en_orden = detalle_existente.cantidad if detalle_existente else 0
        if (cantidad + cantidad_en_orden > stock_disponible):
            return None, f"La cantidad total excede el stock disponible ({stock_disponible})."
        
        if detalle_existente:
            detalle_existente.cantidad += cantidad
            if detalle_existente.cantidad <= 0:
                detalle_existente.delete()
            else:
                if descuento_linea is not None:
                    detalle_existente.descuento = Decimal(descuento_linea)
                detalle_existente.save()
        elif cantidad > 0:
            Detallefactura.objects.create(
                factura=factura, producto=producto_base, variante=variante_encontrada,
                cantidad=cantidad, precio_unitario=precio_final,
                descuento=Decimal(descuento_linea or "0.00"),
                subtotal_linea=Decimal("0.00"), iva_linea=Decimal("0.00"), total_linea=Decimal("0.00")
            )

    recalcular_y_guardar_factura(factura)

    return factura, None



class ActualizarDescuentoDetalleView(APIView):
    """
    Aplica un descuento a una línea de detalle específica y recalcula la factura entera.
    """
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        detalle_id = request.data.get("detalle_id")
        descuento_str = request.data.get("descuento")

        if not detalle_id or descuento_str is None:
            return JsonResponse({"estado": "error", "mensaje": "Se requieren detalle_id y descuento."}, status=HTTPStatus.BAD_REQUEST)
        
        try:
            detalle = Detallefactura.objects.get(id=detalle_id)
            detalle.descuento = Decimal(descuento_str)
            detalle.save(update_fields=['descuento'])

            # Llama a la función centralizada para recalcular todo
            recalcular_y_guardar_factura(detalle.factura)

            # Prepara la respuesta completa que el frontend necesita para actualizarse
            factura_actualizada = Factura.objects.get(id=detalle.factura.id)
            detalles = Detallefactura.objects.filter(factura=factura_actualizada).select_related('producto', 'variante')

            data_detalles = []
            for d in detalles:
                 data_detalles.append({
                    'id': d.id,
                    'producto_id': d.producto.id,
                    'nombre': f"{d.producto.nombre} ({d.variante.nombre})" if d.variante else d.producto.nombre,
                    'codigo_barras': d.variante.codigo_barras if d.variante else d.producto.codigo_barras,
                    'cantidad': d.cantidad,
                    'precio_unitario': str(d.precio_unitario),
                    'descuento': str(d.descuento),
                    'subtotal_linea': str(d.subtotal_linea),
                    'iva_linea': str(d.iva_linea),
                    'total_linea': str(d.total_linea),
                })

            data_factura = {
                'id': factura_actualizada.id,
                'subtotal': str(factura_actualizada.subtotal),
                'iva_total': str(factura_actualizada.iva_total),
                'total': str(factura_actualizada.total),
                'detalles': data_detalles # Devolvemos la lista completa de detalles
            }

            return JsonResponse({ "estado": "success", "factura": data_factura })

        except Detallefactura.DoesNotExist:
             return JsonResponse({"estado": "error", "mensaje": "El detalle no existe."}, status=HTTPStatus.NOT_FOUND)
        except Exception as e:
            return JsonResponse({"estado": "error", "mensaje": str(e)}, status=HTTPStatus.INTERNAL_SERVER_ERROR)



class ActualizarDescuentoGlobalView(APIView):
    """
    Aplica un descuento global a TODAS las líneas de una factura y recalcula totales.
    """
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        factura_id = request.data.get("factura_id")
        descuento_global_str = request.data.get("descuento_global")

        if not factura_id or descuento_global_str is None:
            return JsonResponse({"estado": "error", "mensaje": "Se requieren factura_id y descuento_global."}, status=HTTPStatus.BAD_REQUEST)

        try:
            factura = Factura.objects.get(id=factura_id)
            descuento_global = Decimal(descuento_global_str)

            # 1. Aplicar el descuento a cada detalle y a la factura
            Detallefactura.objects.filter(factura=factura).update(descuento=descuento_global)
            factura.descuento_global = descuento_global
            factura.save(update_fields=['descuento_global'])

            # 2. LLAMAR A LA FUNCIÓN CENTRALIZADA PARA RECALCULAR TODO
            recalcular_y_guardar_factura(factura)

            # 3. Preparar la respuesta completa para el frontend
            factura.refresh_from_db()
            detalles_actualizados = Detallefactura.objects.filter(factura=factura).select_related('producto', 'variante')

            # --- ESTE ES EL BLOQUE QUE FALTABA ---
            data_detalles = []
            for d in detalles_actualizados:
                if d.variante:
                    nombre_item, codigo_barras_item = f"{d.producto.nombre} ({d.variante.nombre})", d.variante.codigo_barras
                else:
                    nombre_item, codigo_barras_item = d.producto.nombre, d.producto.codigo_barras

                data_detalles.append({
                    "id": d.id, "producto_id": d.producto.id, "nombre": nombre_item,
                    "codigo_barras": codigo_barras_item, "cantidad": d.cantidad,
                    "precio_unitario": str(d.precio_unitario), "descuento": str(d.descuento),
                    "subtotal_linea": str(d.subtotal_linea), "iva_linea": str(d.iva_linea),
                    "total_linea": str(d.total_linea),
                })

            data_factura = {
                "id": factura.id, "subtotal": str(factura.subtotal), "iva_total": str(factura.iva_total),
                "total": str(factura.total), "descuento_global": str(factura.descuento_global),
                "detalles": data_detalles
            }
            # --- FIN DEL BLOQUE QUE FALTABA ---
            
            return JsonResponse({"estado": "success", "factura": data_factura})

        except Factura.DoesNotExist:
            return JsonResponse({"estado": "error", "mensaje": "La factura no existe."}, status=HTTPStatus.NOT_FOUND)
        except Exception as e:
            return JsonResponse({"estado": "error", "mensaje": str(e)}, status=HTTPStatus.INTERNAL_SERVER_ERROR)


class FacturaPendienteView(APIView):
    """
    Recupera la factura 'abierta' de un cliente para el día actual.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        cliente_id = request.query_params.get("cliente_id")
        if not cliente_id:
            return JsonResponse({"estado": "error", "mensaje": "cliente_id es requerido."}, status=HTTPStatus.BAD_REQUEST)
        
        factura = Factura.objects.filter(
            cliente_id=cliente_id,
            estado="abierta",
            fecha_operacion__date=timezone.now().date()
        ).first()

        if factura:
            # --- LÓGICA CRÍTICA DE CONSTRUCCIÓN DE RESPUESTA ---
            detalles_queryset = Detallefactura.objects.filter(factura=factura).select_related("producto", "variante")
            detalles_para_frontend = []
            for d in detalles_queryset:
                if d.variante:
                    nombre_item, codigo_barras_item = f"{d.producto.nombre} ({d.variante.nombre})", d.variante.codigo_barras
                else:
                    nombre_item, codigo_barras_item = d.producto.nombre, d.producto.codigo_barras
                
                detalles_para_frontend.append({
                    "id": d.id, "producto_id": d.producto.id, "nombre": nombre_item,
                    "codigo_barras": codigo_barras_item, "cantidad": d.cantidad,
                    "precio_unitario": str(d.precio_unitario), "descuento": str(d.descuento),
                    "subtotal_linea": str(d.subtotal_linea), "iva_linea": str(d.iva_linea),
                    "total_linea": str(d.total_linea)
                })
            
            data_factura = {
                "id": factura.id, "subtotal": str(factura.subtotal or 0),
                "iva_total": str(factura.iva_total or 0), "total": str(factura.total or 0),
                "detalles": detalles_para_frontend,
                "fecha_operacion": factura.fecha_operacion.strftime("%Y-%m-%d %H:%M:%S"),
            }
            return JsonResponse({"estado": "success", "factura": data_factura})
        else:
            return JsonResponse({"estado": "error", "mensaje": "No hay factura pendiente para hoy."}, status=HTTPStatus.NOT_FOUND)


class ResetFacturaView(APIView):
    """
    Busca una factura 'abierta' y la cambia a 'cancelada', borrando sus detalles.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request, *args, **kwargs):
        cliente_id = request.data.get("cliente_id")
        if not cliente_id:
            return JsonResponse({"estado": "error", "mensaje": "cliente_id es requerido."}, status=HTTPStatus.BAD_REQUEST)
        
        factura_abierta = Factura.objects.filter(
            cliente_id=cliente_id,
            estado="abierta",
            fecha_operacion__date=timezone.now().date()
        ).first()

        if factura_abierta:
            # Borra los detalles primero
            Detallefactura.objects.filter(factura=factura_abierta).delete()
            # Luego cambia el estado de la factura a cancelada
            factura_abierta.estado = "cancelada"
            factura_abierta.save()

            return JsonResponse(
                {"estado": "success", "mensaje": "Factura pendiente cancelada. Puede iniciar una nueva orden."},
                status=HTTPStatus.OK
            )
        else:
            return JsonResponse(
                {"estado": "info", "mensaje": "No hay factura pendiente para reiniciar. Puede iniciar una nueva orden."},
                status=HTTPStatus.OK 
            )




# Endpoint que utiliza la función actualizada
class BarcodeScanView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, *args, **kwargs):
        cliente_id = request.data.get("cliente_id")
        barcode = request.data.get("codigo_barras")
        cantidad_str = request.data.get("cantidad", "1")
        eliminar = request.data.get("eliminar", False)
        
        try:
            cantidad = int(cantidad_str)
        except (ValueError, TypeError):
            cantidad = 1

        if not cliente_id or not barcode:
            return JsonResponse(
                {"estado": "error", "mensaje": "Faltan datos de cliente o código de barras."},
                status=HTTPStatus.BAD_REQUEST,
            )

        # --- CAMBIO CLAVE: Se pasa el objeto 'request' a la función ---
        factura, error = agg_producto_a_orden(request, cliente_id, barcode, cantidad, eliminar=eliminar)
        
        if error:
            return JsonResponse(
                {"estado": "error", "mensaje": error},
                status=HTTPStatus.NOT_FOUND,
            )

        # La lógica para construir la respuesta JSON se mantiene igual
        detalles_queryset = Detallefactura.objects.filter(factura=factura).select_related("producto", "variante")
        detalles_para_frontend = []
        for d in detalles_queryset:
            if d.variante:
                nombre_item = f"{d.producto.nombre} ({d.variante.nombre})"
                codigo_barras_item = d.variante.codigo_barras
            else:
                nombre_item = d.producto.nombre
                codigo_barras_item = d.producto.codigo_barras
            
            detalles_para_frontend.append({
                "id": d.id,
                "producto_id": d.producto.id,
                "nombre": nombre_item,
                "codigo_barras": codigo_barras_item,
                "cantidad": d.cantidad,
                "precio_unitario": str(d.precio_unitario),
                "descuento": str(d.descuento),
                "subtotal_linea": str(d.subtotal_linea),
                "iva_linea": str(d.iva_linea),
                "total_linea": str(d.total_linea)
            })

        return JsonResponse({
            "estado": "success",
            "mensaje": "Producto agregado/actualizado en la factura.",
            "factura": {
                "id": factura.id,
                "subtotal": str(factura.subtotal),
                "iva_total": str(factura.iva_total),
                "total": str(factura.total),
                "detalles": detalles_para_frontend
            },
        }, status=HTTPStatus.OK)

# =========================================================
# PRODUCTO CATEGORIA
# =========================================================
class ProductocategoriaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Productocategoria.objects.get(id=id)
            serializer = ProductocategoriaSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Productocategoria.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ProductoCategoria no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class ProductocategoriaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Productocategoria.objects.all().order_by("id")
        serializer = ProductocategoriaSerializer(queryset, many=True)
        return JsonResponse(
            {"producto_categoria": serializer.data}, status=HTTPStatus.OK
        )


class ProductocategoriaCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = ProductocategoriaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Productocategoria.objects.get(id=id)
        except Productocategoria.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ProductoCategoria no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = ProductocategoriaSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Productocategoria.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "ProductoCategoria eliminada"},
                status=HTTPStatus.OK,
            )
        except Productocategoria.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ProductoCategoria no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# =========================================================
# PROVEEDOR
# =========================================================
class ProveedorGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Proveedor.objects.get(id=id)
            serializer = ProveedorSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Proveedor.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Proveedor no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class ProveedorList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Proveedor.objects.all().order_by("id")
        serializer = ProveedorSerializer(queryset, many=True)
        return JsonResponse({"proveedores": serializer.data}, status=HTTPStatus.OK)


class ProveedorCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = ProveedorSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Proveedor.objects.get(id=id)
        except Proveedor.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Proveedor no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = ProveedorSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Proveedor.objects.get(id=id, activo=True)
            obj.activo = False
            obj.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Proveedor eliminado"},
                status=HTTPStatus.OK,
            )
        except Proveedor.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Proveedor no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# =========================================================
# REPORTES (Solo lectura) y EXPORTACIÓN CSV
# =========================================================


def generar_pdf(queryset, headers, title):
    """Función genérica para crear un PDF a partir de un queryset."""
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{title.lower().replace(" ", "_")}.pdf"'
    
    p = canvas.Canvas(response, pagesize=landscape(letter))
    width, height = landscape(letter)
    
    # Título
    p.setFont("Helvetica-Bold", 16)
    p.drawString(inch, height - inch, title)
    
    # Cabeceras de la tabla
    p.setFont("Helvetica-Bold", 10)
    x = inch
    y = height - 1.5 * inch
    col_width = (width - 2 * inch) / len(headers)
    
    for header in headers:
        p.drawString(x, y, header)
        x += col_width
    y -= 15
    
    # Contenido de la tabla
    p.setFont("Helvetica", 9)
    for item in queryset:
        x = inch
        for header in headers:
            # getattr anida para obtener valores de ForeignKeys (ej: 'cliente.nombre')
            value = item
            for part in header.lower().replace(" ", "_").split('.'):
                value = getattr(value, part, '')
            p.drawString(x, y, str(value or ''))
            x += col_width
        y -= 15
        if y < inch: # Salto de página
            p.showPage()
            p.setFont("Helvetica", 9)
            y = height - inch
    
    p.showPage()
    p.save()
    return response

def generar_excel(queryset, headers, title):
    """Función genérica para crear un Excel a partir de un queryset."""
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{title.lower().replace(" ", "_")}.xlsx"'
    
    wb = Workbook()
    ws = wb.active
    ws.title = title
    
    # Estilo para cabeceras
    header_font = Font(bold=True)
    
    # Escribir cabeceras
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
    
    # Escribir datos
    for row_num, item in enumerate(queryset, 2):
        for col_num, header in enumerate(headers, 1):
            value = item
            for part in header.lower().replace(" ", "_").split('.'):
                 value = getattr(value, part, '')
            ws.cell(row=row_num, column=col_num, value=str(value or ''))
            
    wb.save(response)
    return response

# --- VISTAS UNIFICADAS PARA REPORTES ---

class ReporteclienteView(APIView):
    """
    Vista para listar y exportar el reporte de clientes.
    Soporta JSON, PDF, Excel y CSV a través del parámetro '?export=[pdf|excel|csv]'
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        export_format = request.query_params.get('export', None)
        queryset = Reportecliente.objects.all().order_by("cliente__nombre")
        headers = ["Cliente", "Total Compras", "Cantidad Pedidos", "Producto Mas Comprado"]
        
        if export_format == 'pdf':
            return generar_pdf(queryset, headers, "Reporte de Clientes")
            
        if export_format == 'excel':
            return generar_excel(queryset, headers, "Reporte de Clientes")

        if export_format == 'csv':
            response = HttpResponse(content_type="text/csv")
            response["Content-Disposition"] = 'attachment; filename="reporte_clientes.csv"'
            writer = csv.writer(response)
            writer.writerow(headers)
            for reporte in queryset:
                writer.writerow([
                    reporte.cliente,
                    reporte.total_compras,
                    reporte.cantidad_pedidos,
                    reporte.producto_mas_comprado,
                ])
            return response
            
        # Por defecto, devuelve JSON
        serializer = ReporteclienteSerializer(queryset, many=True)
        return JsonResponse({"reportes": serializer.data}, status=HTTPStatus.OK)


class InventarioActualView(APIView):
    """
    Vista para obtener un reporte de inventario en tiempo real,
    con opción de filtrar por categoría.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        categoria_id = request.query_params.get('categoria_id', None)

        # Obtener productos simples y variantes que deben estar en el inventario
        productos_simples = Producto.objects.filter(activo=True, tipo='simple')
        variantes = Variacionproducto.objects.select_related('producto__categoria').filter(producto__activo=True)

        if categoria_id:
            try:
                # Filtrar ambos querysets por el ID de la categoría
                productos_simples = productos_simples.filter(categoria_id=categoria_id)
                variantes = variantes.filter(producto__categoria_id=categoria_id)
            except ValueError:
                return Response({"error": "ID de categoría inválido."}, status=400)

        # Unificar los datos en una sola lista con una estructura consistente
        inventario_unificado = []
        for producto in productos_simples:
            inventario_unificado.append({
                'id': f"p-{producto.id}",
                'nombre': producto.nombre,
                'categoria': producto.categoria.nombre if producto.categoria else 'Sin Categoría',
                'codigo_barras': producto.codigo_barras,
                'cantidad': producto.cantidad or 0
            })
        
        for variante in variantes:
            inventario_unificado.append({
                'id': f"v-{variante.id}",
                'nombre': f"{variante.producto.nombre} - {variante.nombre}",
                'categoria': variante.producto.categoria.nombre if variante.producto.categoria else 'Sin Categoría',
                'codigo_barras': variante.codigo_barras,
                'cantidad': variante.cantidad or 0
            })

        # Ordenar la lista final alfabéticamente por nombre
        inventario_ordenado = sorted(inventario_unificado, key=lambda item: item['nombre'])
        
        # Serializar los datos usando el serializer genérico
        serializer = InventarioItemSerializer(inventario_ordenado, many=True)
        
        # Obtener todas las categorías para el filtro del frontend (solo se hace una vez)
        categorias_data = []
        if not categoria_id: # Para no enviarlas en cada filtro
             categorias = Categoriaproducto.objects.filter(activo=True).order_by('nombre')
             categorias_data = CategoriaSerializer(categorias, many=True).data

        return Response({
            "inventario": serializer.data,
            "categorias": categorias_data
        })


class ReporteventaView(APIView):
    """
    Genera un reporte detallado de ventas para el frontend.
    Esta vista devuelve una lista de facturas individuales usando un serializer optimizado
    que no incluye los detalles de los productos.
    """
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        start_date_str = request.query_params.get("start_date")
        end_date_str = request.query_params.get("end_date")
        estado = request.query_params.get('estado', None)

        if not start_date_str or not end_date_str:
            return Response({"error": "Se requieren start_date y end_date"}, status=400)

        try:
            # Convertir las fechas de string a objeto datetime consciente de la zona horaria
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
            
            start_of_day = timezone.make_aware(datetime.combine(start_date, time.min))
            end_of_day = timezone.make_aware(datetime.combine(end_date, time.max))

            # Construir el queryset base, optimizando con select_related
            queryset = Factura.objects.filter(
                fecha_operacion__range=(start_of_day, end_of_day)
            ).select_related('cliente', 'usuario', 'metodo_pago').order_by('-fecha_operacion')

            # Aplicar filtro de estado si se proporciona
            if estado:
                queryset = queryset.filter(estado__iexact=estado)
            
            # --- CAMBIO CLAVE ---
            # Usar el VentaReporteSerializer para convertir los datos al formato JSON correcto
            serializer = VentaReporteSerializer(queryset, many=True)
            
            # Devolver los datos en el formato que el frontend espera ("reporte_detallado")
            return Response({"reporte_detallado": serializer.data}, status=200)

        except Exception as e:
            print(f"Error generando reporte de ventas: {e}")
            return Response({"error": "Ocurrió un error inesperado."}, status=500)



# =========================================================
# RESERVA STOCK
# =========================================================
class ReservastockGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Reservastock.objects.get(id=id)
            serializer = ReservastockSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Reservastock.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ReservaStock no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class ReservastockList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Reservastock.objects.all().order_by("id")
        serializer = ReservastockSerializer(queryset, many=True)
        return JsonResponse({"reserva_stock": serializer.data}, status=HTTPStatus.OK)


class ReservastockCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = ReservastockSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Reservastock.objects.get(id=id)
        except Reservastock.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ReservaStock no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = ReservastockSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Reservastock.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "ReservaStock eliminada"},
                status=HTTPStatus.OK,
            )
        except Reservastock.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ReservaStock no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# =========================================================
# ROL
# =========================================================
class RolGet(APIView):
    permission_classes = [IsAuthenticated,IsAdmin]
    def get(self, request, id):
        try:
            obj = Rol.objects.get(id=id)
            serializer = RolSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Rol.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Rol no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class RolList(APIView):
    permission_classes = [IsAuthenticated,IsAdmin]
    def get(self, request):
        queryset = Rol.objects.all().order_by("id")
        serializer = RolSerializer(queryset, many=True)
        return Response({"roles": serializer.data}, status=status.HTTP_200_OK)


class RolCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated,IsAdmin]
    def post(self, request):
        serializer = RolSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Rol.objects.get(id=id)
        except Rol.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Rol no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = RolSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Rol.objects.get(id=id, activo=True)
            obj.activo = False
            obj.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Rol eliminado"},
                status=HTTPStatus.OK,
            )
        except Rol.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Rol no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# =========================================================
# SESION USUARIO
# =========================================================
class SesionusuarioGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Sesionusuario.objects.get(id=id)
            serializer = SesionusuarioSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Sesionusuario.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "SesionUsuario no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class SesionusuarioList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Sesionusuario.objects.all().order_by("id")
        serializer = SesionusuarioSerializer(queryset, many=True)
        return JsonResponse({"sesiones": serializer.data}, status=HTTPStatus.OK)


class SesionusuarioCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated,IsAdmin]
    def post(self, request):
        serializer = SesionusuarioSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creada", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Sesionusuario.objects.get(id=id)
        except Sesionusuario.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "SesionUsuario no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = SesionusuarioSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizada", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Sesionusuario.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminada", "mensaje": "SesionUsuario eliminada"},
                status=HTTPStatus.OK,
            )
        except Sesionusuario.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "SesionUsuario no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# =========================================================
# TIPO DOCUMENTO FISCAL
# =========================================================
class TipodocumentofiscalGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Tipodocumentofiscal.objects.get(id=id)
            serializer = TipodocumentofiscalSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Tipodocumentofiscal.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TipoDocumentoFiscal no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class TipodocumentofiscalList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Tipodocumentofiscal.objects.all().order_by("id")
        serializer = TipodocumentofiscalSerializer(queryset, many=True)
        return JsonResponse({"tipos_documento": serializer.data}, status=HTTPStatus.OK)


class TipodocumentofiscalCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = TipodocumentofiscalSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Tipodocumentofiscal.objects.get(id=id)
        except Tipodocumentofiscal.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TipoDocumentoFiscal no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = TipodocumentofiscalSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Tipodocumentofiscal.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "TipoDocumentoFiscal eliminado"},
                status=HTTPStatus.OK,
            )
        except Tipodocumentofiscal.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TipoDocumentoFiscal no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# ============
#Metodo Pago

class MetodoPagoList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = MetodoPago.objects.filter(activo=True).order_by("nombre")
        serializer = MetodoPagoSerializer(queryset, many=True)
        return JsonResponse({"metodos": serializer.data}, status=HTTPStatus.OK)


class MetodoPagoCRUD(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = MetodoPago.objects.get(id=id)
            serializer = MetodoPagoSerializer(data)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except MetodoPago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Metodo de pago no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


    def post(self, request):
        serializer = MetodoPagoSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = MetodoPago.objects.get(id=id)
        except MetodoPago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Método de pago no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = MetodoPagoSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = MetodoPago.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Método de pago eliminado"},
                status=HTTPStatus.OK,
            )
        except MetodoPago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Método de pago no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )



# =========================================================
# TRANSACCION PAGO
# =========================================================
def reducir_stock_de_factura(factura):
    """
    Recorre los detalles de una factura, reduce el stock local y, si el item
    tiene SKU, envía la actualización a WooCommerce.
    """
    detalles = Detallefactura.objects.select_related('producto', 'variante').filter(factura=factura)
    
    for detalle in detalles:
        item_a_actualizar = None
        
        if detalle.variante:
            # CASO 1: Es una variante.
            variante = detalle.variante
            stock_actual = variante.cantidad or 0
            variante.cantidad = stock_actual - detalle.cantidad
            variante.save(update_fields=['cantidad'])
            item_a_actualizar = variante 
            
        elif detalle.producto:
            # CASO 2: Es un producto simple.
            producto = detalle.producto
            stock_actual = producto.cantidad or 0
            producto.cantidad = stock_actual - detalle.cantidad
            producto.save(update_fields=['cantidad'])
            item_a_actualizar = producto 

        # --- 2. LLAMA A LA FUNCIÓN DE SINCRONIZACIÓN ---
        # Después de guardar el cambio local, actualizamos WooCommerce.
        if item_a_actualizar and hasattr(item_a_actualizar, 'sku') and item_a_actualizar.sku:
            print(f"Iniciando sincronización para SKU: {item_a_actualizar.sku}")
            actualizar_stock_woocommerce(
                sku=item_a_actualizar.sku, 
                nueva_cantidad=item_a_actualizar.cantidad
            )


class PagoView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        factura_id = request.data.get("factura_id")
        monto_recibido_str = request.data.get("monto_recibido")
        metodo_pago_id = request.data.get("metodo_pago")
        datos_adicionales = request.data.get("datos_adicionales", {})
        estado_override = request.data.get("estado", "").lower()

        if not factura_id or not metodo_pago_id:
            return JsonResponse({"estado": "error", "mensaje": "Se requiere factura_id y método de pago."}, status=HTTPStatus.BAD_REQUEST)

        try:
            metodo = MetodoPago.objects.get(id=metodo_pago_id)
            factura = Factura.objects.select_for_update().get(id=factura_id, estado="abierta")
        except MetodoPago.DoesNotExist:
            return JsonResponse({"estado": "error", "mensaje": "Método de pago no encontrado."}, status=HTTPStatus.NOT_FOUND)
        except Factura.DoesNotExist:
            return JsonResponse({"estado": "error", "mensaje": "Factura no encontrada o ya procesada."}, status=HTTPStatus.NOT_FOUND)
        
        # --- Lógica de "Pagar luego" o estado "pendiente" ---
        if metodo.tipo_metodo == "Pagar luego" or estado_override == "pendiente":
            if not factura.correlativo:
                # Lógica para asignar correlativo
                pass

            # LLAMADA A LA FUNCIÓN AUXILIAR PARA REDUCIR STOCK (ahora también sincroniza)
            reducir_stock_de_factura(factura)

            factura.metodo_pago = metodo
            factura.estado = "pendiente"
            factura.nombre_cliente_pendiente = datos_adicionales.get("nombre_cliente", "")
            factura.comentario_pendiente = datos_adicionales.get("comentario", "")
            factura.save()

            return JsonResponse({"estado": "success", "mensaje": "La factura ha sido marcada como pendiente."}, status=HTTPStatus.OK)
        
        # --- Lógica para pagos completos ---
        if monto_recibido_str is None or float(monto_recibido_str) < float(factura.total):
            return JsonResponse({"estado": "error", "mensaje": f"Monto insuficiente. Se requiere {factura.total}"}, status=HTTPStatus.BAD_REQUEST)
        
        if not factura.correlativo:
            # Lógica para asignar correlativo
            pass

        # LLAMADA A LA FUNCIÓN AUXILIAR PARA REDUCIR STOCK (ahora también sincroniza)
        reducir_stock_de_factura(factura)

        factura.metodo_pago = metodo
        factura.estado = "pagado"
        factura.fecha_pago = timezone.now()
        factura.save()

        # Crear la transacción de pago
        codigo_transaccion = str(uuid.uuid4())
        transaccion_data = {
            "factura": factura.id, 
            "monto": factura.total,
            "metodo_pago": metodo_pago_id,
            "estado": "exitoso",
            "codigo_transaccion": codigo_transaccion,
            "fecha": timezone.now(),
            "activo": True,
        }
        
        serializer = TransaccionpagoSerializer(data=transaccion_data)
        if serializer.is_valid():
            serializer.save()
        else:
            return JsonResponse({"estado": "error", "mensaje": "Error al registrar la transacción.", "detalles": serializer.errors}, status=HTTPStatus.BAD_REQUEST)

        return JsonResponse({
            "estado": "success",
            "mensaje": "Pago completado y transacción registrada.",
            "factura": {"id": factura.id, "estado": factura.estado, "total": str(factura.total)},
            "transaccion": serializer.data
        }, status=HTTPStatus.OK)

class FacturaImprimirView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, factura_id, *args, **kwargs):
        try:
            # Obtener el tamaño del papel desde los parámetros de la URL (default: 80)
            paper_size = request.GET.get('paper_size', '80')

            # --- Lógica para seleccionar los archivos CSS ---
            base_css_path = os.path.join(settings.BASE_DIR, 'erp', 'static', 'css', 'ticket_base.css')
            
            if paper_size == '58':
                size_css_path = os.path.join(settings.BASE_DIR, 'erp', 'static', 'css', 'ticket_58mm.css')
            else: # Default a 80mm
                size_css_path = os.path.join(settings.BASE_DIR, 'erp', 'static', 'css', 'ticket_80mm.css')

            # Cargar los datos de la factura (sin cambios)
            factura = Factura.objects.select_related('cliente', 'usuario', 'metodo_pago').get(id=factura_id)
            detalles = factura.detallefactura_set.all()
            nombre_empleado = factura.usuario.get_full_name().strip() or factura.usuario.username if factura.usuario else "N/A"

            logo_base64 = None
            logo_path = os.path.join(settings.BASE_DIR, 'img', 'LOGO-POWER.png')
            if os.path.exists(logo_path):
                with open(logo_path, "rb") as image_file:
                    encoded_string = base64.b64encode(image_file.read()).decode('utf-8')
                    logo_base64 = f"data:image/png;base64,{encoded_string}"

            context = {
                'factura': factura,
                'detalles': detalles,
                'nombre_empleado': nombre_empleado,
                'empresa': {
                    'nombre': "Power Nutricion Deportiva Excelente, S.L.",
                    'cif': "B-55450688",
                    'direccion': "Calle miguel de prado, 4 BJ 47002, Valladolid",
                    'telefono': "641 00 89 57",
                    'logo_path': logo_base64
                }
            }

            html_string = render_to_string('tickets/ticket_template.html', context)
            
            # --- Generar PDF cargando AMBOS archivos CSS ---
            stylesheets = [CSS(base_css_path), CSS(size_css_path)]
            pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri()).write_pdf(stylesheets=stylesheets)

            response = HttpResponse(pdf_file, content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="factura_{factura.correlativo}.pdf"'
            
            return response

        except Factura.DoesNotExist:
            return JsonResponse({"estado": "error", "mensaje": "La factura no existe."}, status=404)
        except Exception as e:
            print(f"Error al generar el PDF para factura {factura_id}: {e}")
            return JsonResponse({"estado": "error", "mensaje": "Error al generar la factura."}, status=500)


class TransaccionpagoGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            obj = Transaccionpago.objects.get(id=id)
            serializer = TransaccionpagoSerializer(obj)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Transaccionpago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TransaccionPago no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class TransaccionpagoList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        queryset = Transaccionpago.objects.all().order_by("id")
        serializer = TransaccionpagoSerializer(queryset, many=True)
        return JsonResponse({"transacciones": serializer.data}, status=HTTPStatus.OK)


class TransaccionpagoCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = TransaccionpagoSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creada", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            obj = Transaccionpago.objects.get(id=id)
        except Transaccionpago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TransaccionPago no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = TransaccionpagoSerializer(obj, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizada", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            obj = Transaccionpago.objects.get(id=id)
            obj.delete()
            return JsonResponse(
                {"estado": "eliminada", "mensaje": "TransaccionPago eliminada"},
                status=HTTPStatus.OK,
            )
        except Transaccionpago.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "TransaccionPago no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


#Filtrar Transacciones por Método de Pago
class TransaccionpagoByMetodo(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, metodo):
        queryset = Transaccionpago.objects.filter(metodo_pago=metodo).order_by("-fecha")
        serializer = TransaccionpagoSerializer(queryset, many=True)
        return JsonResponse({"transacciones": serializer.data}, status=HTTPStatus.OK)




# =========================================================
# VARIACION PRODUCTO, ATRIBUTOS Y VALOR
# =========================================================

# class ValorAtributoListView(APIView):
#     permission_classes = [IsAuthenticated]
#     def get(self, request):
#         valores = ValorAtributo.objects.all()
#         serializer = ValorAtributoSerializer(valores, many=True)
#         return Response(serializer.data)

class AtributoListCreateView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        # Ahora prefetch_related('valores') funcionará gracias al related_name
        atributos = Atributo.objects.prefetch_related('valores').all()
        serializer = AtributoSerializer(atributos, many=True)
        return Response(serializer.data) # <-- CORREGIDO: Usar Response

    def post(self, request):
        serializer = AtributoSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class AtributoDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def get_object(self, pk):
        try:
            return Atributo.objects.get(pk=pk)
        except Atributo.DoesNotExist:
            raise Http404

    def put(self, request, pk):
        atributo = self.get_object(pk)
        serializer = AtributoSerializer(atributo, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data) # <-- CORREGIDO: Usar Response
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def delete(self, request, pk):
        atributo = self.get_object(pk)
        atributo.delete()
        return Response(status=status.HTTP_204_NO_CONTENT) # <-- CORREGIDO: Usar Response
    
class VariacionproductoListCreateView(ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    queryset = Variacionproducto.objects.all()
    serializer_class = VariacionproductoSerializer

class VariacionproductoDetailView(RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    queryset = Variacionproducto.objects.all()
    serializer_class = VariacionproductoSerializer

# ==============================================================
# Configuración IVA
# ==============================================================
class ConfiguracionivaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            config = Configuracioniva.objects.get(id=id)
            serializer = ConfiguracionivaSerializer(config)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Configuracioniva.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ConfiguracionIVA no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class ConfiguracionivaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        configs = Configuracioniva.objects.all().order_by("id")
        serializer = ConfiguracionivaSerializer(configs, many=True)
        return JsonResponse({"configuraciones": serializer.data}, status=HTTPStatus.OK)


class ConfiguracionivaCRUD(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = ConfiguracionivaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            config = Configuracioniva.objects.get(id=id)
        except Configuracioniva.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ConfiguracionIVA no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = ConfiguracionivaSerializer(config, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            config = Configuracioniva.objects.get(id=id)
            config.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "ConfiguracionIVA eliminada"},
                status=HTTPStatus.OK,
            )
        except Configuracioniva.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "ConfiguracionIVA no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Cupón Descuento
# ==============================================================
class CupondescuentoGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            cupon = Cupondescuento.objects.get(id=id)
            serializer = CupondescuentoSerializer(cupon)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Cupondescuento.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Cupón de descuento no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class CupondescuentoList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        cupones = Cupondescuento.objects.all().order_by("id")
        serializer = CupondescuentoSerializer(cupones, many=True)
        return JsonResponse({"cupones": serializer.data}, status=HTTPStatus.OK)


class CupondescuentoCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = CupondescuentoSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            cupon = Cupondescuento.objects.get(id=id)
        except Cupondescuento.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Cupón de descuento no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = CupondescuentoSerializer(cupon, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            cupon = Cupondescuento.objects.get(id=id)
            cupon.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Cupón de descuento eliminado"},
                status=HTTPStatus.OK,
            )
        except Cupondescuento.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Cupón de descuento no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Detalle Factura
# ==============================================================
class DetallefacturaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            detalle = Detallefactura.objects.get(id=id)
            serializer = DetallefacturaSerializer(detalle)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Detallefactura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Detalle de factura no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class DetallefacturaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        detalles = Detallefactura.objects.all().order_by("id")
        serializer = DetallefacturaSerializer(detalles, many=True)
        return JsonResponse({"detalles": serializer.data}, status=HTTPStatus.OK)


class DetallefacturaCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = DetallefacturaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            detalle = Detallefactura.objects.get(id=id)
        except Detallefactura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Detalle de factura no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = DetallefacturaSerializer(detalle, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            detalle = Detallefactura.objects.get(id=id)
            detalle.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Detalle de factura eliminado"},
                status=HTTPStatus.OK,
            )
        except Detallefactura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Detalle de factura no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Devolución
# ==============================================================
class DevolucionGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            devolucion = Devolucion.objects.get(id=id)
            serializer = DevolucionSerializer(devolucion)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Devolucion.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Devolución no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class DevolucionList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        devoluciones = Devolucion.objects.all().order_by("id")
        serializer = DevolucionSerializer(devoluciones, many=True)
        return JsonResponse({"devoluciones": serializer.data}, status=HTTPStatus.OK)


class DevolucionCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = DevolucionSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            devolucion = Devolucion.objects.get(id=id)
        except Devolucion.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Devolución no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = DevolucionSerializer(devolucion, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            devolucion = Devolucion.objects.get(id=id, activo=True)
            devolucion.activo = False
            devolucion.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Devolución eliminada"},
                status=HTTPStatus.OK,
            )
        except Devolucion.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Devolución no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Factura
# ==============================================================
class FacturaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            factura = Factura.objects.get(id=id)
            serializer = FacturaReporteSerializer(factura)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Factura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class FacturaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        facturas = Factura.objects.all().order_by("id")
        serializer = FacturaSerializer(facturas, many=True)
        return JsonResponse({"facturas": serializer.data}, status=HTTPStatus.OK)


class FacturaCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = FacturaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            factura = Factura.objects.get(id=id)
        except Factura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = FacturaSerializer(factura, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            factura = Factura.objects.get(id=id, activo=True)
            factura.activo = False
            factura.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Factura eliminada"},
                status=HTTPStatus.OK,
            )
        except Factura.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Factura Electrónica
# ==============================================================
class FacturaelectronicaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            fact_elec = Facturaelectronica.objects.get(id=id)
            serializer = FacturaelectronicaSerializer(fact_elec)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Facturaelectronica.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura electrónica no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class FacturaelectronicaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        fact_elems = Facturaelectronica.objects.all().order_by("id")
        serializer = FacturaelectronicaSerializer(fact_elems, many=True)
        return JsonResponse(
            {"facturas_electronicas": serializer.data}, status=HTTPStatus.OK
        )


class FacturaelectronicaCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = FacturaelectronicaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            fact_elec = Facturaelectronica.objects.get(id=id)
        except Facturaelectronica.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura electrónica no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = FacturaelectronicaSerializer(
            fact_elec, data=request.data, partial=True
        )
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            fact_elec = Facturaelectronica.objects.get(id=id, activo=True)
            fact_elec.activo = False
            fact_elec.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Factura electrónica eliminada"},
                status=HTTPStatus.OK,
            )
        except Facturaelectronica.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Factura electrónica no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Movimiento Inventario
# ==============================================================
class MovimientoInventarioGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            movimiento = MovimientoInventario.objects.get(id=id)
            serializer = MovimientoInventarioSerializer(movimiento)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except MovimientoInventario.DoesNotExist:
            return JsonResponse(
                {
                    "estado": "error",
                    "mensaje": "Movimiento de inventario no encontrado",
                },
                status=HTTPStatus.NOT_FOUND,
            )


class MovimientoInventarioList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        movimientos = MovimientoInventario.objects.all().order_by("id")
        serializer = MovimientoInventarioSerializer(movimientos, many=True)
        return JsonResponse({"movimientos": serializer.data}, status=HTTPStatus.OK)


class MovimientoInventarioCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = MovimientoInventarioSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            movimiento = MovimientoInventario.objects.get(id=id)
        except MovimientoInventario.DoesNotExist:
            return JsonResponse(
                {
                    "estado": "error",
                    "mensaje": "Movimiento de inventario no encontrado",
                },
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = MovimientoInventarioSerializer(
            movimiento, data=request.data, partial=True
        )
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            movimiento = MovimientoInventario.objects.get(id=id, estado=True)
            movimiento.activo = False
            movimiento.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Movimiento de inventario inactivo"},
                status=HTTPStatus.OK,
            )
        except MovimientoInventario.DoesNotExist:
            return JsonResponse(
                {
                    "estado": "error",
                    "mensaje": "Movimiento de inventario no encontrado",
                },
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Lectura Código Barras
# ==============================================================
class LecturacodigobarrasGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            lectura = Lecturacodigobarras.objects.get(id=id)
            serializer = LecturacodigobarrasSerializer(lectura)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Lecturacodigobarras.DoesNotExist:
            return JsonResponse(
                {
                    "estado": "error",
                    "mensaje": "Lectura de código de barras no encontrada",
                },
                status=HTTPStatus.NOT_FOUND,
            )


class LecturacodigobarrasList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        lecturas = Lecturacodigobarras.objects.all().order_by("id")
        serializer = LecturacodigobarrasSerializer(lecturas, many=True)
        return JsonResponse({"lecturas": serializer.data}, status=HTTPStatus.OK)


class LecturacodigobarrasCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = LecturacodigobarrasSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            lectura = Lecturacodigobarras.objects.get(id=id)
        except Lecturacodigobarras.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Lectura no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = LecturacodigobarrasSerializer(
            lectura, data=request.data, partial=True
        )
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            lectura = Lecturacodigobarras.objects.get(id=id)
            lectura.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Lectura eliminada"},
                status=HTTPStatus.OK,
            )
        except Lecturacodigobarras.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Lectura no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Log de Actividad
# ==============================================================
class LogactividadGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            log = Logactividad.objects.get(id=id)
            serializer = LogactividadSerializer(log)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Logactividad.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Log de actividad no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class LogactividadList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        logs = Logactividad.objects.all().order_by("id")
        serializer = LogactividadSerializer(logs, many=True)
        return JsonResponse({"logs": serializer.data}, status=HTTPStatus.OK)


class LogactividadCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = LogactividadSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            log = Logactividad.objects.get(id=id)
        except Logactividad.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Log de actividad no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = LogactividadSerializer(log, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            log = Logactividad.objects.get(id=id)
            log.delete()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Log de actividad eliminado"},
                status=HTTPStatus.OK,
            )
        except Logactividad.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Log de actividad no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==============================================================
# Orden
# ==============================================================
class OrdenGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            orden = Orden.objects.get(id=id)
            serializer = OrdenSerializer(orden)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Orden.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Orden no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class OrdenList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        ordenes = Orden.objects.all().order_by("id")
        serializer = OrdenSerializer(ordenes, many=True)
        return JsonResponse({"ordenes": serializer.data}, status=HTTPStatus.OK)


class OrdenCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = OrdenSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def put(self, request, id):
        try:
            orden = Orden.objects.get(id=id)
        except Orden.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Orden no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )
        serializer = OrdenSerializer(orden, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            orden = Orden.objects.get(id=id, activo=True)
            orden.activo = False
            orden.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Orden eliminada"},
                status=HTTPStatus.OK,
            )
        except Orden.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Orden no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


# ==========
# Almacen
# ==========


class AlmacenGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = Almacen.objects.get(id=id)
            serializer = AlmacenSerializer(data)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Exception as e:
            return JsonResponse(
                {"estado": "error", "mensaje": "Almacen no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class AlmacenList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        almacen = Almacen.objects.filter(activo=True).order_by("id")
        serializer = AlmacenSerializer(almacen, many=True)

        return JsonResponse({"almacenes": serializer.data}, status=HTTPStatus.OK)


class AlmacenCRUD(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = AlmacenSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def put(self, request, id):
        try:
            almacen = Almacen.objects.get(id=id)
        except Almacen.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Almacen no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )

        serializer = AlmacenSerializer(almacen, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def delete(self, request, id):
        try:
            almacen = Almacen.objects.get(id=id, activo=True)
            almacen.activo = False
            almacen.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "almacen inactivo"},
                status=HTTPStatus.OK,
            )
        except Producto.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Almacen no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class InventarioGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = Inventario.objects.get(id=id)
            serializer = InventarioSerializer(data)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Exception as e:
            return JsonResponse(
                {"estado": "error", "mensaje": "Inventario no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class InventarioList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        data = Inventario.objects.filter(activo=True).order_by("id")
        serializer = InventarioSerializer(data, many=True)
        return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)


class InventarioCrud(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = InventarioSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def put(self, request, id):
        try:
            data = Inventario.objects.get(id=id)
        except Inventario.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Inventario no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )

        serializer = InventarioSerializer(data, request=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )

        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            data = Inventario.objects.get(id=id, activo=True)
            data.activo = False
            data.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Inventario inactivo"},
                status=HTTPStatus.OK,
            )
        except Inventario.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Inventario no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class CategoriaGet(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = Categoriaproducto.objects.get(id=id)
            serializer = CategoriaSerializer(data)
            return JsonResponse({"data": serializer.data}, status=HTTPStatus.OK)
        except Exception as e:
            return JsonResponse(
                {"estado": "error", "mensaje": "Categoria no encontrada"},
                status=HTTPStatus.NOT_FOUND,
            )


class CategoriaList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        categoria = Categoriaproducto.objects.filter(activo=True).order_by("id")
        serializer = CategoriaSerializer(categoria, many=True)
        return JsonResponse({"categorias": serializer.data}, status=HTTPStatus.OK)


class CategoriaCrud(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = CategoriaSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save(creado_por=request.user)
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def put(self, request, id):
        try:
            categoria = Categoriaproducto.objects.get(id=id)
        except Categoriaproducto.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Categoria no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )

        serializer = CategoriaSerializer(categoria, data=request.data, partial=True)

        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )

        return JsonResponse(
            {"estado": "error", "mensaje": serializer.errors},
            status=HTTPStatus.BAD_REQUEST,
        )

    def delete(self, request, id):
        try:
            categoria = Categoriaproducto.objects.get(id=id, activo=True)
            categoria.activo = False
            categoria.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Categoria inactiva"},
                status=HTTPStatus.OK,
            )
        except Categoriaproducto.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Categoria no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )


class ObtenerclienteAPIView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, id):
        try:
            data = Cliente.objects.get(id=id)
            serializer = ClienteSerializer(data)
            return JsonResponse(
                {"estado": "ok", "mensaje": serializer.data}, status=status.HTTP_200_OK
            )
        except Cliente.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=status.HTTP_404_NOT_FOUND,
            )


class ClienteList(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        clientes = Cliente.objects.filter(activo=True).order_by("id")
        serializer = ClienteSerializer(clientes, many=True)

        return Response({"clientes": serializer.data}, status=HTTPStatus.OK)

# --- 2. Vista para el Reporte de Facturas (la lista principal) ---
class FacturaDetalleReporteView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        start_date_str = request.query_params.get("start_date")
        end_date_str = request.query_params.get("end_date")
        estado = request.query_params.get('estado', None)
        cliente_id = request.query_params.get('cliente_id', None)

        try:
            # Convertir las fechas de string a objeto date
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()

            # --- LÓGICA DE ZONA HORARIA CORREGIDA ---
            # Crear un rango de 24 horas para cada fecha, consciente de la zona horaria
            start_of_day = timezone.make_aware(datetime.combine(start_date, time.min))
            end_of_day = timezone.make_aware(datetime.combine(end_date, time.max))

            # Filtramos usando el rango consciente de la zona horaria
            queryset = Factura.objects.filter(
                fecha_operacion__range=(start_of_day, end_of_day)
            ).select_related('cliente', 'usuario').prefetch_related('detallefactura_set').order_by('-fecha_operacion')

            if estado:
                queryset = queryset.filter(estado__iexact=estado) # Usamos iexact para ignorar mayúsculas/minúsculas
            if cliente_id:
                queryset = queryset.filter(cliente_id=cliente_id)
            
            serializer = FacturaReporteSerializer(queryset, many=True)
            return Response({"reporte_detallado": serializer.data}, status=status.HTTP_200_OK)

        except (ValueError, TypeError):
            return Response({"error": "Formato de fecha inválido. Se requiere YYYY-MM-DD."}, status=status.HTTP_400_BAD_REQUEST)

class ClienteCreateUpdateDelete(APIView):
    permission_classes = [IsAuthenticated]
    def post(self, request):
        serializer = ClienteSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "creado", "data": serializer.data}, status=HTTPStatus.CREATED
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def put(self, request, id):
        try:
            cliente = Cliente.objects.get(id=id)
            serializer = ClienteSerializer(cliente, data=request.data, partial=True)
        except Cliente.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Cliente no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )

        if serializer.is_valid():
            serializer.save()
            return JsonResponse(
                {"estado": "actualizado", "data": serializer.data}, status=HTTPStatus.OK
            )
        else:
            return JsonResponse(
                {"estado": "error", "mensaje": serializer.errors},
                status=HTTPStatus.BAD_REQUEST,
            )

    def delete(self, request, id):
        try:
            cliente = Cliente.objects.get(id=id, activo=True)
            cliente.activo = False
            cliente.save()
            return JsonResponse(
                {"estado": "eliminado", "mensaje": "Cliente inactivo"},
                status=HTTPStatus.OK,
            )
        except Cliente.DoesNotExist:
            return JsonResponse(
                {"estado": "error", "mensaje": "Cliente no encontrado"},
                status=HTTPStatus.NOT_FOUND,
            )

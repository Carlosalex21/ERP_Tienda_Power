# apps/commons/utils.py
from django.http import HttpResponse
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font

# CLASE UTILITARIA

def generar_pdf(queryset, headers, title):
    """Función genérica para crear un PDF a partir de un queryset."""
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{title.lower().replace(" ", "_")}.pdf"'
    
    p = canvas.Canvas(response, pagesize=landscape(letter))
    width, height = landscape(letter)
    
    # Título
    p.setFont("Helvetica-Bold", 16)
    p.drawString(inch, height - inch, title)
    
    # Cabeceras de la tabla
    p.setFont("Helvetica-Bold", 10)
    x = inch
    y = height - 1.5 * inch
    col_width = (width - 2 * inch) / len(headers)
    
    for header in headers:
        p.drawString(x, y, header)
        x += col_width
    y -= 15
    
    # Contenido de la tabla
    p.setFont("Helvetica", 9)
    for item in queryset:
        x = inch
        for header in headers:
            value = item
            for part in header.lower().replace(" ", "_").split('.'):
                value = getattr(value, part, '')
            p.drawString(x, y, str(value or ''))
            x += col_width
        y -= 15
        if y < inch: # Salto de página
            p.showPage()
            p.setFont("Helvetica", 9)
            y = height - inch
    
    p.showPage()
    p.save()
    return response


def generar_excel(queryset, headers, title):
    """Función genérica para crear un Excel a partir de un queryset."""
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{title.lower().replace(" ", "_")}.xlsx"'
    
    wb = Workbook()
    ws = wb.active
    ws.title = title
    
    # Estilo para cabeceras
    header_font = Font(bold=True)
    
    # Escribir cabeceras
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
    
    # Escribir datos
    for row_num, item in enumerate(queryset, 2):
        for col_num, header in enumerate(headers, 1):
            value = item
            for part in header.lower().replace(" ", "_").split('.'):
                value = getattr(value, part, '')
            ws.cell(row=row_num, column=col_num, value=str(value or ''))
            
    wb.save(response)
    return response
"""Exportación de los reportes de Contabilidad a PDF (xhtml2pdf) y Excel (openpyxl)."""
from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from xhtml2pdf import pisa
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import EmpresaContable


def _nombre_tenant() -> str:
    try:
        from apps.configuracion.models import ConfiguracionEmpresa
        config, _ = ConfiguracionEmpresa.objects.get_or_create(pk=1)
        return config.nombre_comercial or config.razon_social or ''
    except Exception:
        return ''


def _pdf_response(template_name: str, context: dict, nombre_archivo: str, download: bool) -> HttpResponse:
    context = {
        **context,
        'fecha_generacion': timezone.now().strftime('%d/%m/%Y %H:%M'),
        'tenant_nombre': _nombre_tenant(),
    }
    html = render_to_string(template_name, context)
    result = BytesIO()
    pisa.CreatePDF(html, dest=result)
    response = HttpResponse(result.getvalue(), content_type='application/pdf')
    disposicion = 'attachment' if download else 'inline'
    response['Content-Disposition'] = f'{disposicion}; filename="{nombre_archivo}.pdf"'
    return response


def _excel_response(nombre_archivo: str, hojas: list[tuple[str, list[str], list[list]]]) -> HttpResponse:
    """`hojas` es una lista de (titulo_hoja, encabezados, filas)."""
    wb = Workbook()
    wb.remove(wb.active)
    encabezado_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    encabezado_font = Font(color='FFFFFF', bold=True)

    for titulo, encabezados, filas in hojas:
        ws = wb.create_sheet(title=titulo[:31])
        ws.append(encabezados)
        for celda in ws[1]:
            celda.fill = encabezado_fill
            celda.font = encabezado_font
        for fila in filas:
            ws.append(fila)
        for columna in ws.columns:
            longitud = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
            ws.column_dimensions[columna[0].column_letter].width = min(longitud + 3, 45)

    result = BytesIO()
    wb.save(result)
    response = HttpResponse(
        result.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.xlsx"'
    return response


# --- Balance de Comprobación ---

def balance_comprobacion_pdf(empresa: EmpresaContable, filas: list[dict], fecha_desde, fecha_hasta, download: bool) -> HttpResponse:
    from decimal import Decimal
    total_debe = sum((Decimal(f['debe']) for f in filas), Decimal('0.00'))
    total_haber = sum((Decimal(f['haber']) for f in filas), Decimal('0.00'))
    context = {
        'empresa': empresa, 'filas': filas, 'fecha_desde': fecha_desde, 'fecha_hasta': fecha_hasta,
        'total_debe': f'{total_debe:.2f}', 'total_haber': f'{total_haber:.2f}',
        'cuadra': total_debe == total_haber,
    }
    return _pdf_response('contabilidad/balance_comprobacion_pdf.html', context, f'balance_comprobacion_{empresa.id}', download)


def balance_comprobacion_excel(empresa: EmpresaContable, filas: list[dict]) -> HttpResponse:
    encabezados = ['Código', 'Cuenta', 'Debe', 'Haber', 'Saldo']
    datos = [[f['codigo'], f['nombre'], float(f['debe']), float(f['haber']), float(f['saldo'])] for f in filas]
    return _excel_response(f'balance_comprobacion_{empresa.id}', [('Balance de Comprobación', encabezados, datos)])


# --- Estados Financieros ---

def estados_financieros_pdf(empresa: EmpresaContable, datos: dict, fecha_desde, fecha_hasta, download: bool) -> HttpResponse:
    context = {
        'empresa': empresa, 'bg': datos['balance_general'], 'er': datos['estado_resultados'],
        'fecha_desde': fecha_desde, 'fecha_hasta': fecha_hasta,
    }
    return _pdf_response('contabilidad/estados_financieros_pdf.html', context, f'estados_financieros_{empresa.id}', download)


def estados_financieros_excel(empresa: EmpresaContable, datos: dict) -> HttpResponse:
    bg = datos['balance_general']
    er = datos['estado_resultados']
    encabezados = ['Cuenta', 'Saldo']

    filas_bg = [['ACTIVO', '']] + [[f['nombre'], float(f['saldo'])] for f in bg['activo']] + [['Total Activo', float(bg['total_activo'])]]
    filas_bg += [['PASIVO', '']] + [[f['nombre'], float(f['saldo'])] for f in bg['pasivo']] + [['Total Pasivo', float(bg['total_pasivo'])]]
    filas_bg += [['PATRIMONIO', '']] + [[f['nombre'], float(f['saldo'])] for f in bg['patrimonio']] + [['Total Patrimonio', float(bg['total_patrimonio'])]]
    filas_bg += [['Utilidad del período', float(bg['utilidad_periodo'])]]

    filas_er = [['INGRESOS', '']] + [[f['nombre'], float(f['saldo'])] for f in er['ingresos']] + [['Total Ingresos', float(er['total_ingresos'])]]
    filas_er += [['COSTOS', '']] + [[f['nombre'], float(f['saldo'])] for f in er['costos']] + [['Total Costos', float(er['total_costos'])]]
    filas_er += [['GASTOS', '']] + [[f['nombre'], float(f['saldo'])] for f in er['gastos']] + [['Total Gastos', float(er['total_gastos'])]]
    filas_er += [['Utilidad del Período', float(er['utilidad_periodo'])]]

    return _excel_response(f'estados_financieros_{empresa.id}', [
        ('Balance General', encabezados, filas_bg),
        ('Estado de Resultados', encabezados, filas_er),
    ])


# --- Libro Mayor ---

def libro_mayor_pdf(empresa: EmpresaContable, cuenta, datos: dict, fecha_desde, fecha_hasta, download: bool) -> HttpResponse:
    context = {
        'empresa': empresa, 'cuenta': cuenta, 'movimientos': datos['movimientos'],
        'saldo_final': datos['saldo_final'], 'fecha_desde': fecha_desde, 'fecha_hasta': fecha_hasta,
    }
    return _pdf_response('contabilidad/libro_mayor_pdf.html', context, f'libro_mayor_{cuenta.id}', download)


def libro_mayor_excel(empresa: EmpresaContable, cuenta, datos: dict) -> HttpResponse:
    encabezados = ['Fecha', 'Asiento', 'Descripción', 'Debe', 'Haber', 'Saldo']
    filas = [
        [str(m['fecha']), f"#{m['asiento_numero']}", m['descripcion'], float(m['debe']), float(m['haber']), float(m['saldo'])]
        for m in datos['movimientos']
    ]
    return _excel_response(f'libro_mayor_{cuenta.id}', [(f'{cuenta.codigo} {cuenta.nombre}', encabezados, filas)])
